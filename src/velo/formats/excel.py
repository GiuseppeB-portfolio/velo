"""Excel (.xlsx) adapter.

The first row of every sheet is the header. A field id is 'Sheet!Header'
(duplicated headers in the same sheet become 'Header#2', ...). Ids are never
parsed, only compared with the ones computed from the file, so sheet names or
headers containing '!' are fine.

Cells holding formulas are never touched. Cells not in the plan are never
rewritten, so their type, number format and style stay as they were.
"""

from __future__ import annotations

import re
import warnings as warnings_module
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from ..errors import UnsupportedFileError
from ..model import ApplyReport, FieldInfo, Inspection, Plan, Replacer
from .base import SAMPLE_SIZE, FormatAdapter, check_plan, preview_text, unique_labels

_SHAPE_RE = re.compile(rb"<(?:\w+:)?sp[ >]")
_METADATA_FIELDS = (
    "creator",
    "lastModifiedBy",
    "title",
    "subject",
    "description",
    "keywords",
    "category",
)


def _open(path: Path):
    """Load a workbook, returning it with the warnings openpyxl emitted.

    openpyxl warns when it drops something it does not understand
    (extensions, unsupported drawing parts...): those warnings are
    information the user must see.
    """
    if Path(path).suffix.lower() != ".xlsx":
        raise UnsupportedFileError(
            "Solo i file .xlsx sono supportati (.xls e .xlsm no: "
            "le macro e il vecchio formato andrebbero persi)."
        )
    with warnings_module.catch_warnings(record=True) as caught:
        warnings_module.simplefilter("always")
        try:
            wb = openpyxl.load_workbook(path)
        except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
            raise UnsupportedFileError("Il file non è un .xlsx valido.") from exc
    return wb, [str(w.message) for w in caught]


def _scan(ws) -> list[dict[str, Any]]:
    """Describe the columns of a sheet (header in row 1, data from row 2)."""
    cols = []
    for col in range(1, ws.max_column + 1):
        header = ws.cell(1, col).value
        label = str(header).strip() if header is not None else ""
        samples: list[str] = []
        non_empty = formulas = 0
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, col)
            if cell.data_type == "f":
                formulas += 1
                continue
            if cell.value is None or cell.value == "":
                continue
            non_empty += 1
            if len(samples) < SAMPLE_SIZE:
                samples.append(preview_text(cell.value))
        if not label:
            if not (non_empty or formulas):
                continue  # entirely empty, unnamed column
            label = f"Colonna {get_column_letter(col)}"
        cols.append(dict(col=col, label=label, sample=tuple(samples), non_empty=non_empty, formulas=formulas))
    for item, uniq in zip(cols, unique_labels(c["label"] for c in cols)):
        item["label"] = uniq
        item["id"] = f"{ws.title}!{uniq}"
    return cols


def _package_warnings(path: Path, wb) -> list[str]:
    out: list[str] = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        if any(n.startswith("xl/externalLinks/") for n in names):
            out.append(
                "Il file collega altri file Excel: dopo la scrittura i valori "
                "memorizzati di quelle formule vanno persi."
            )
        if any(n.startswith("xl/pivotTables/") for n in names):
            out.append("Ci sono tabelle pivot: potrebbero non essere conservate fedelmente.")
        if any(n.startswith("xl/slicers/") for n in names):
            out.append("Ci sono filtri interattivi (slicer): andranno persi.")
        if any(n.startswith("xl/threadedComments/") for n in names):
            out.append("Ci sono commenti a catena: andranno persi e restano nel file originale.")
        if any(n.startswith("xl/media/") for n in names):
            try:
                import PIL  # noqa: F401
            except ImportError:
                out.append("Ci sono immagini ma Pillow non è installato: andranno perse.")
            else:
                out.append("Ci sono immagini: controlla che l'output le mostri come l'originale.")
        for n in names:
            if re.fullmatch(r"xl/drawings/drawing\d+\.xml", n) and _SHAPE_RE.search(z.read(n)):
                out.append("Ci sono forme o caselle di testo: andranno perse.")
                break
        if any(n.startswith("xl/comments") for n in names):
            out.append("Ci sono commenti nelle celle: restano in chiaro (anche autore e testo).")
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            out.append(f"Il foglio '{ws.title}' è nascosto ma i suoi campi sono elencati.")
    return out


class ExcelAdapter(FormatAdapter):
    name = "excel"
    extensions = (".xlsx",)

    def inspect(self, path: Path) -> Inspection:
        wb, load_warnings = _open(path)
        fields: list[FieldInfo] = []
        for ws in wb.worksheets:
            for c in _scan(ws):
                fields.append(
                    FieldInfo(
                        id=c["id"],
                        label=c["label"],
                        container=ws.title,
                        sample=c["sample"],
                        non_empty=c["non_empty"],
                        formula_cells=c["formulas"],
                    )
                )
        warns = load_warnings + _package_warnings(Path(path), wb)
        warns.append(
            "Nomi dei fogli, intestazioni, nomi definiti e testo scritto dentro le "
            "formule non vengono modificati."
        )
        return Inspection("excel", fields, {"scrub_metadata": True}, warns)

    def apply(
        self,
        src: Path,
        dst: Path,
        plan: Plan,
        replacer: Replacer,
        params: dict[str, Any] | None = None,
    ) -> ApplyReport:
        p = {"scrub_metadata": True, **(params or {})}
        wb, load_warnings = _open(src)
        scans = {ws.title: _scan(ws) for ws in wb.worksheets}
        check_plan(plan, (c["id"] for cols in scans.values() for c in cols))

        replaced = {fid: 0 for fid in plan}
        skipped = {fid: 0 for fid in plan}
        for ws in wb.worksheets:
            for c in scans[ws.title]:
                choice = plan.get(c["id"])
                if choice is None:
                    continue
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row, c["col"])
                    if isinstance(cell, MergedCell):
                        continue  # only the top-left cell of a merge holds a value
                    if cell.data_type == "f":
                        skipped[c["id"]] += 1
                        continue
                    if cell.value is None or cell.value == "":
                        continue
                    new = replacer(choice, cell.value)
                    cell.value = new
                    if isinstance(new, str) and new.startswith("="):
                        cell.data_type = "s"  # a text must never become a formula
                    replaced[c["id"]] += 1

        warns = load_warnings + _package_warnings(Path(src), wb)
        for fid, n in replaced.items():
            if n == 0:
                reason = " (contiene solo formule, che non vengono toccate)" if skipped[fid] else ""
                warns.append(f"Il campo '{fid}' non conteneva valori da sostituire{reason}.")
        if p["scrub_metadata"]:
            for name in _METADATA_FIELDS:
                setattr(wb.properties, name, None)
            wb.properties.creator = "velo"  # openpyxl would otherwise write its own name
        wb.save(dst)
        return ApplyReport(replaced=replaced, skipped_formulas=skipped, warnings=warns)
