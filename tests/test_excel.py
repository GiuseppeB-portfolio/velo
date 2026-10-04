from datetime import datetime

import openpyxl
import pytest

from factories import identity, make_workbook, tag
from velo.errors import PlanError, UnsupportedFileError
from velo.formats.excel import ExcelAdapter
from velo.model import Category, FieldChoice, Mode

A = ExcelAdapter()


def choice(fid, cat=Category.TEXT):
    return FieldChoice(fid, cat, Mode.FAKE)


@pytest.fixture
def wbpath(tmp_path):
    return make_workbook(tmp_path / "in.xlsx")


def test_inspect_lists_all_sheets_with_preview(wbpath):
    insp = A.inspect(wbpath)
    ids = [f.id for f in insp.fields]
    assert ids == [
        "Clienti!Nome",
        "Clienti!Cognome",
        "Clienti!Importo",
        "Clienti!Data",
        "Clienti!IVA",
        "Clienti!Completo",
        "Clienti!Note",
        "Ordini!Cliente",
        "Ordini!Nome",
        "Ordini!Quantita",
        "Ordini!Nome#2",
        "Archivio!Nome",
    ]
    nome = insp.fields[0]
    assert nome.sample == ("Mario", "Lucia", "Anna", "Ufficio")
    iva = insp.fields[4]
    assert iva.formula_cells == 3 and iva.non_empty == 0 and iva.sample == ()
    assert any("'Archivio' è nascosto" in w for w in insp.warnings)


def test_identity_keeps_everything(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"
    A.apply(wbpath, dst, {}, identity, params={"scrub_metadata": False})
    a, b = openpyxl.load_workbook(wbpath), openpyxl.load_workbook(dst)
    assert a.sheetnames == b.sheetnames
    for name in a.sheetnames:
        wa, wb_ = a[name], b[name]
        assert wa.sheet_state == wb_.sheet_state
        assert wa.merged_cells.ranges == wb_.merged_cells.ranges
        for row in wa.iter_rows():
            for ca in row:
                cb = wb_[ca.coordinate]
                assert (ca.value, ca.data_type, ca.number_format, ca.font.b) == (
                    cb.value,
                    cb.data_type,
                    cb.number_format,
                    cb.font.b,
                )


def test_chosen_fields_change_in_every_sheet_others_do_not(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"
    plan = {
        "Clienti!Nome": choice("Clienti!Nome"),
        "Ordini!Nome": choice("Ordini!Nome"),
        "Ordini!Nome#2": choice("Ordini!Nome#2"),
    }
    rep = A.apply(wbpath, dst, plan, tag)
    wb = openpyxl.load_workbook(dst)
    c, o = wb["Clienti"], wb["Ordini"]
    assert [c[f"A{r}"].value for r in (2, 3, 4, 5)] == ["X[Mario]", "X[Lucia]", "X[Anna]", "X[Ufficio]"]
    assert [o[f"B{r}"].value for r in (2, 3)] == ["X[Mario]", "X[Lucia]"]
    assert [o[f"D{r}"].value for r in (2, 3)] == ["X[Rossi]", "X[Bianchi]"]
    assert c["B2"].value == "Rossi" and o["A2"].value == "C1"  # untouched
    assert wb["Archivio"]["A2"].value == "Segreto"  # not chosen -> intact
    assert rep.replaced == {"Clienti!Nome": 4, "Ordini!Nome": 2, "Ordini!Nome#2": 2}


def test_types_formats_and_layout_of_untouched_cells(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"
    A.apply(wbpath, dst, {"Clienti!Nome": choice("Clienti!Nome")}, tag)
    c = openpyxl.load_workbook(dst)["Clienti"]
    assert isinstance(c["C2"].value, float) and c["C2"].number_format == "#,##0.00"
    assert isinstance(c["C4"].value, int)
    assert c["D2"].value == datetime(1980, 2, 1) and c["D2"].number_format == "DD/MM/YYYY"
    assert c["A1"].font.b is True
    assert c.column_dimensions["A"].width == 22
    assert c.freeze_panes == "A2"
    assert "A5:A6" in {str(r) for r in c.merged_cells.ranges}
    assert len(c.data_validations.dataValidation) == 1


def test_formulas_are_never_touched_even_if_chosen(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"
    plan = {"Clienti!IVA": choice("Clienti!IVA"), "Clienti!Completo": choice("Clienti!Completo")}
    rep = A.apply(wbpath, dst, plan, tag)
    c = openpyxl.load_workbook(dst)["Clienti"]
    assert c["E2"].value == "=C2*0.22" and c["F3"].value == '=A3&" "&B3'
    assert c["F2"].data_type == "f"
    assert rep.replaced == {"Clienti!IVA": 0, "Clienti!Completo": 0}
    assert rep.skipped_formulas == {"Clienti!IVA": 3, "Clienti!Completo": 3}
    assert any("solo formule" in w for w in rep.warnings)


def test_replacement_types_are_kept(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"

    def fake(ch, v):
        return datetime(2000, 1, 1) if isinstance(v, datetime) else 1.0

    A.apply(
        wbpath,
        dst,
        {"Clienti!Data": choice("Clienti!Data"), "Clienti!Importo": choice("Clienti!Importo")},
        fake,
    )
    c = openpyxl.load_workbook(dst)["Clienti"]
    assert c["D2"].value == datetime(2000, 1, 1) and c["D2"].number_format == "DD/MM/YYYY"
    assert c["C2"].value == 1.0 and c["C2"].number_format == "#,##0.00"


def test_text_starting_with_equals_never_becomes_a_formula(wbpath, tmp_path):
    dst = tmp_path / "out.xlsx"
    A.apply(wbpath, dst, {"Clienti!Cognome": choice("Clienti!Cognome")}, lambda c, v: "=1+1")
    cell = openpyxl.load_workbook(dst)["Clienti"]["B2"]
    assert cell.data_type == "s" and cell.value == "=1+1"


def _docprops(path):
    import zipfile

    with zipfile.ZipFile(path) as z:
        return b"".join(z.read(n) for n in z.namelist() if n.startswith("docProps/"))


def test_metadata_is_scrubbed_by_default_and_can_be_kept(wbpath, tmp_path):
    assert b"Giuseppe" in _docprops(wbpath) and b"Acme" in _docprops(wbpath)  # the fixture has them
    A.apply(wbpath, tmp_path / "a.xlsx", {}, identity)
    raw = _docprops(tmp_path / "a.xlsx")
    assert b"Giuseppe" not in raw and b"Acme" not in raw
    A.apply(wbpath, tmp_path / "b.xlsx", {}, identity, params={"scrub_metadata": False})
    assert b"Giuseppe" in _docprops(tmp_path / "b.xlsx")


def test_unknown_field_is_an_error(wbpath, tmp_path):
    with pytest.raises(PlanError, match="Clienti!Nomi"):
        A.apply(wbpath, tmp_path / "o.xlsx", {"Clienti!Nomi": choice("Clienti!Nomi")}, tag)
    assert not (tmp_path / "o.xlsx").exists()


@pytest.mark.parametrize("name", ["a.xls", "a.xlsm", "a.csv"])
def test_other_extensions_are_refused(tmp_path, name):
    p = tmp_path / name
    p.write_bytes(b"x")
    with pytest.raises(UnsupportedFileError):
        A.inspect(p)


def test_corrupt_xlsx_is_refused(tmp_path):
    p = tmp_path / "a.xlsx"
    p.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedFileError):
        A.inspect(p)


def test_losses_are_reported_when_writing_too(tmp_path):
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    wb.active.append(["Nome"])
    wb.active.append(["Mario"])
    wb.active["A2"].comment = Comment("telefono 333 1234567", "Giuseppe")
    wb.save(tmp_path / "c.xlsx")
    assert any("commenti" in w for w in A.inspect(tmp_path / "c.xlsx").warnings)
    rep = A.apply(tmp_path / "c.xlsx", tmp_path / "o.xlsx", {"Sheet!Nome": choice("Sheet!Nome")}, tag)
    assert any("commenti" in w for w in rep.warnings)
