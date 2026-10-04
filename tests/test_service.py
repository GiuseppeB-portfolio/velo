import json
import stat
import sys
from datetime import datetime

import openpyxl
import pytest

from factories import CSV_IT, make_workbook, write_csv
from velo.errors import PlanError, VeloError
from velo.keystore import default_dir, ephemeral_key, key_id, load_or_create_project_key
from velo.mapping import MappingFile
from velo.model import Category as C
from velo.model import FieldChoice, Mode
from velo.service import pseudonymize, restore
from velo.validators import is_valid_fiscal_code


def plan_of(*pairs, mode=Mode.FAKE):
    return {fid: FieldChoice(fid, cat, mode) for fid, cat in pairs}


# -- CSV -----------------------------------------------------------------------------------

CSV_PLAN = plan_of(
    ("Nome", C.FIRST_NAME),
    ("Cognome", C.LAST_NAME),
    ("Importo", C.AMOUNT),
    ("Citta", C.CITY),
    ("Data", C.DATE),
)


def test_csv_roundtrip_restores_the_original_bytes(tmp_path):
    src = write_csv(tmp_path / "clienti.csv", CSV_IT, "cp1252")
    r = pseudonymize(src, tmp_path / "out", CSV_PLAN)
    out = r.output.read_bytes()
    for secret in ("Mario", "Rossi", "Bianchi", "Lucia", "Torino", "01/02/1980", "1.234,56"):
        assert secret.encode("cp1252") not in out
    back = restore(r.output, r.mapping, tmp_path / "back.csv")
    assert (tmp_path / "back.csv").read_bytes() == src.read_bytes()
    assert back.identical_to_original is True and not back.unknown


def test_csv_untouched_columns_and_dialect(tmp_path):
    src = write_csv(
        tmp_path / "c.csv", "Nome;Nota;Imp\r\nMario;ciao;1.234,56\r\nLucia;ciao;9,9\r\n", "cp1252"
    )
    r = pseudonymize(src, tmp_path / "out", plan_of(("Nome", C.FIRST_NAME)))
    lines = r.output.read_bytes().decode("cp1252").split("\r\n")
    assert lines[0] == "Nome;Nota;Imp"
    assert lines[1].endswith(";ciao;1.234,56") and lines[2].endswith(";ciao;9,9")


def test_same_value_gets_same_fake_throughout_the_file(tmp_path):
    src = write_csv(tmp_path / "c.csv", "Nome;Altro\nMario;Mario\nLucia;Mario\nMario;x\n")
    r = pseudonymize(src, tmp_path / "out", plan_of(("Nome", C.FIRST_NAME), ("Altro", C.FIRST_NAME)))
    rows = [l.split(";") for l in r.output.read_text().splitlines()[1:]]
    assert rows[0][0] == rows[0][1] == rows[1][1] == rows[2][0] and rows[0][0] != rows[1][0]


def test_placeholders_in_csv(tmp_path):
    src = write_csv(tmp_path / "c.csv", "Nome;Cognome\nMario;Rossi\n")
    r = pseudonymize(
        src,
        tmp_path / "out",
        plan_of(("Nome", C.FIRST_NAME), ("Cognome", C.LAST_NAME), mode=Mode.PLACEHOLDER),
    )
    first, last = r.output.read_text().splitlines()[1].split(";")
    assert first.startswith("[FIRST_NAME_") and last.startswith("[LAST_NAME_")
    restore(r.output, r.mapping, tmp_path / "b.csv")
    assert (tmp_path / "b.csv").read_bytes() == src.read_bytes()


# -- JSON ----------------------------------------------------------------------------------

JSON_RAW = (
    '{\n  "clienti": [\n    {"anagrafica": {"nome": "Mario Rossi", "citta": "Forl\\u00ec", "cf": "RSSMRA80A01H501U"},\n'
    '     "importo": 1234.50, "nascita": "1980-02-01", "tel": "+39 333 1234567", "ok": true, "x": null},\n'
    '    {"anagrafica": {"nome": "Lucia Bianchi", "citta": "Forl\\u00ec", "cf": ""},\n'
    '     "importo": 99.90, "nascita": "1975-11-15", "tel": "0543 123456", "ok": false, "x": null}\n  ]\n}\n'
)
JSON_PLAN = plan_of(
    ("clienti[].anagrafica.nome", C.FULLNAME),
    ("clienti[].anagrafica.citta", C.CITY),
    ("clienti[].anagrafica.cf", C.FISCAL_CODE),
    ("clienti[].importo", C.AMOUNT),
    ("clienti[].nascita", C.DATE),
    ("clienti[].tel", C.PHONE),
)


def test_json_roundtrip_restores_the_original_bytes(tmp_path):
    src = tmp_path / "d.json"
    src.write_bytes(JSON_RAW.encode())
    r = pseudonymize(src, tmp_path / "out", JSON_PLAN)
    out = json.loads(r.output.read_text())
    c0, c1 = out["clienti"]
    assert c0["anagrafica"]["citta"] == c1["anagrafica"]["citta"]
    assert c0["anagrafica"]["nome"] != "Mario Rossi" and is_valid_fiscal_code(c0["anagrafica"]["cf"])
    assert c1["anagrafica"]["cf"] == "" and c0["x"] is None and c0["ok"] is True
    assert isinstance(c0["importo"], float) and c0["tel"].startswith("+39 3")
    for secret in ("Mario", "Rossi", "Bianchi", "1980-02-01", "1234.50", "1234567", "RSSMRA"):
        assert secret not in r.output.read_text()
    back = restore(r.output, r.mapping, tmp_path / "back.json")
    assert (tmp_path / "back.json").read_bytes() == src.read_bytes()  # incl. 1234.50 and \u00ec
    assert back.identical_to_original is True


def test_json_numeric_placeholder_roundtrip(tmp_path):
    src = tmp_path / "d.json"
    src.write_bytes(b'{"tel": [3331234567, 3337654321], "n": 1.10}')
    r = pseudonymize(src, tmp_path / "out", plan_of(("tel[]", C.PHONE), mode=Mode.PLACEHOLDER))
    assert r.output.read_text().count("[PHONE_") == 2
    restore(r.output, r.mapping, tmp_path / "b.json")
    assert (tmp_path / "b.json").read_bytes() == src.read_bytes()


# -- Excel ---------------------------------------------------------------------------------


def test_excel_roundtrip_restores_values_and_types(tmp_path):
    src = make_workbook(tmp_path / "x.xlsx")
    plan = plan_of(
        ("Clienti!Nome", C.FIRST_NAME),
        ("Clienti!Cognome", C.LAST_NAME),
        ("Clienti!Importo", C.AMOUNT),
        ("Clienti!Data", C.DATE),
        ("Ordini!Nome", C.FIRST_NAME),
        ("Ordini!Nome#2", C.LAST_NAME),
    )
    r = pseudonymize(src, tmp_path / "out", plan)
    c = openpyxl.load_workbook(r.output)["Clienti"]
    o = openpyxl.load_workbook(r.output)["Ordini"]
    assert (
        c["A2"].value != "Mario" and isinstance(c["C2"].value, float) and isinstance(c["D2"].value, datetime)
    )
    assert c["D2"].number_format == "DD/MM/YYYY" and c["E2"].value == "=C2*0.22"
    assert o["B2"].value == c["A2"].value  # "Mario" is the same fake in every sheet
    back = restore(r.output, r.mapping, tmp_path / "back.xlsx")
    a, b = openpyxl.load_workbook(src), openpyxl.load_workbook(tmp_path / "back.xlsx")
    for name in a.sheetnames:
        for row in a[name].iter_rows():
            for cell in row:
                assert b[name][cell.coordinate].value == cell.value, (name, cell.coordinate)
                assert b[name][cell.coordinate].number_format == cell.number_format
    assert not back.unknown


# -- cross-file coherence and keys -------------------------------------------------------------


def test_project_key_links_files_ephemeral_does_not(tmp_path):
    a = write_csv(tmp_path / "a.csv", "Nome\nMario\nLucia\n")
    b = write_csv(tmp_path / "b.csv", "Nome\nLucia\nMario\n")
    p = plan_of(("Nome", C.FIRST_NAME))
    keys = tmp_path / "keys"
    ra = pseudonymize(a, tmp_path / "o1", p, key_mode="project", key_dir=keys)
    rb = pseudonymize(b, tmp_path / "o2", p, key_mode="project", key_dir=keys)
    fa = ra.output.read_text().split()[1:]
    fb = rb.output.read_text().split()[1:]
    assert fa == list(reversed(fb)) and ra.key_id == rb.key_id
    ea = pseudonymize(a, tmp_path / "o3", p)
    eb = pseudonymize(b, tmp_path / "o4", p)
    assert ea.output.read_text().split()[1:] != list(reversed(eb.output.read_text().split()[1:]))
    assert ea.key_id != eb.key_id


def test_project_key_is_created_once_and_reused(tmp_path):
    k1 = load_or_create_project_key(tmp_path)
    assert load_or_create_project_key(tmp_path) == k1 and len(k1) == 32
    assert key_id(k1) == key_id(k1) != key_id(ephemeral_key())
    if sys.platform != "win32":
        assert stat.S_IMODE((tmp_path / "project.key").stat().st_mode) == 0o600
    (tmp_path / "bad").mkdir()
    (tmp_path / "bad" / "project.key").write_text("abcd")
    with pytest.raises(ValueError):
        load_or_create_project_key(tmp_path / "bad")


def test_default_dir_is_under_appdata_on_windows(monkeypatch, tmp_path):
    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setenv("APPDATA", str(tmp_path))
    assert default_dir() == tmp_path / "velo"


# -- mapping file and safety checks ---------------------------------------------------------------


def test_mapping_file_content_and_warning(tmp_path):
    src = write_csv(tmp_path / "c.csv", CSV_IT, "cp1252")
    r = pseudonymize(src, tmp_path / "out", CSV_PLAN)
    doc = json.loads(r.mapping.read_text(encoding="utf-8"))
    assert r.mapping.name == "c.pseudo.csv.velo-map.json"
    assert "ATTENZIONE" in doc["WARNING"] and "originali in chiaro" in doc["WARNING"]
    assert doc["plan"]["Nome"] == {"category": "first_name", "mode": "fake"}
    assert doc["key"]["mode"] == "ephemeral" and len(doc["key"]["id"]) == 8
    assert doc["source"]["format"] == "csv" and len(doc["source"]["sha256"]) == 64
    assert doc["source_params"]["encoding"] == "cp1252" and doc["faker_version"]
    assert "Mario" in json.dumps(doc["entries"])  # real data is in there: that is the point
    assert MappingFile.load(r.mapping).to_plan()["Importo"].category is C.AMOUNT


def test_never_overwrites_and_never_touches_the_original(tmp_path):
    src = write_csv(tmp_path / "c.csv", CSV_IT, "cp1252")
    before = src.read_bytes()
    r = pseudonymize(src, tmp_path, CSV_PLAN)  # same folder as the original
    assert src.read_bytes() == before and r.output.name == "c.pseudo.csv"
    with pytest.raises(VeloError, match="esiste già"):
        pseudonymize(src, tmp_path, CSV_PLAN)
    pseudonymize(src, tmp_path, CSV_PLAN, overwrite=True)
    with pytest.raises(VeloError, match="coinciderebbe"):
        restore(r.output, r.mapping, r.output)  # dst == pseudonymised file


def test_empty_plan_and_bad_plan_leave_no_files(tmp_path):
    src = write_csv(tmp_path / "c.csv", CSV_IT, "cp1252")
    with pytest.raises(VeloError, match="Nessun campo"):
        pseudonymize(src, tmp_path / "out", {})
    with pytest.raises(PlanError):
        pseudonymize(src, tmp_path / "out", plan_of(("Inesistente", C.TEXT)))
    assert sorted(p.name for p in (tmp_path / "out").iterdir()) == []


def test_restore_reports_values_missing_from_the_dictionary(tmp_path):
    src = write_csv(tmp_path / "c.csv", "Nome\nMario\nLucia\n")
    r = pseudonymize(src, tmp_path / "out", plan_of(("Nome", C.FIRST_NAME)))
    edited = r.output.read_text().replace(r.output.read_text().splitlines()[1], "Qualcuno")
    r.output.write_text(edited)
    back = restore(r.output, r.mapping, tmp_path / "b.csv")
    assert back.unknown == {"Nome": 1} and back.identical_to_original is False
    assert any("non sono nel dizionario" in w for w in back.warnings)


def test_restore_rejects_wrong_format_or_bad_mapping(tmp_path):
    src = write_csv(tmp_path / "c.csv", "Nome\nMario\n")
    r = pseudonymize(src, tmp_path / "out", plan_of(("Nome", C.FIRST_NAME)))
    other = tmp_path / "o.json"
    other.write_text('{"Nome": "x"}')
    with pytest.raises(VeloError, match="formato"):
        restore(other, r.mapping, tmp_path / "x.json")
    bad = tmp_path / "bad.velo-map.json"
    bad.write_text("{}")
    with pytest.raises(VeloError, match="dizionario"):
        restore(r.output, bad, tmp_path / "y.csv")


def test_report_carries_engine_notes(tmp_path):
    src = write_csv(tmp_path / "c.csv", "Data\n01/02/1980\nboh\n")
    r = pseudonymize(src, tmp_path / "out", plan_of(("Data", C.DATE)))
    assert any("segnaposto" in w for w in r.warnings)
    restore(r.output, r.mapping, tmp_path / "b.csv")
    assert (tmp_path / "b.csv").read_bytes() == src.read_bytes()
