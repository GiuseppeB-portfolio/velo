import io
import re

import openpyxl
import pytest

from factories import CSV_IT, make_workbook
from velo.web.app import CATEGORIES, create_app, parse_args

BASE = "http://127.0.0.1:5000"


@pytest.fixture
def app(tmp_path):
    return create_app(port=5000, workdir=tmp_path / "work")


@pytest.fixture
def client(app):
    return app.test_client()


def token(app):
    return app.config["CSRF_TOKEN"]


def upload(client, app, raw: bytes, name: str):
    return client.post(
        "/upload",
        base_url=BASE,
        content_type="multipart/form-data",
        data={"csrf": token(app), "file": (io.BytesIO(raw), name)},
    )


def sid_of(resp):
    return re.search(r"/s/([\w-]+)", resp.headers["Location"]).group(1)


def run(client, app, sid, form):
    return client.post(f"/s/{sid}/run", base_url=BASE, data={"csrf": token(app), **form})


CSV = CSV_IT.encode("cp1252")
CSV_FORM = {
    "sel_0": "on",
    "cat_0": "first_name",
    "mode_0": "fake",
    "sel_1": "on",
    "cat_1": "last_name",
    "mode_1": "fake",
    "sel_3": "on",
    "cat_3": "city",
    "mode_3": "placeholder",
}


# -- protections ----------------------------------------------------------------------------


@pytest.mark.parametrize("host", ["evil.example", "127.0.0.1", "127.0.0.1:9999", "localhost.evil.com:5000"])
def test_wrong_host_is_refused(client, host):
    assert client.get("/", base_url=f"http://{host}").status_code == 400


@pytest.mark.parametrize("host", ["127.0.0.1:5000", "localhost:5000"])
def test_loopback_hosts_are_accepted(client, host):
    assert client.get("/", base_url=f"http://{host}").status_code == 200


def test_post_without_or_with_wrong_token_is_refused(client, app):
    data = {"file": (io.BytesIO(CSV), "a.csv")}
    assert (
        client.post("/upload", base_url=BASE, data=data, content_type="multipart/form-data").status_code
        == 403
    )
    data = {"file": (io.BytesIO(CSV), "a.csv"), "csrf": "x" * 43}
    assert (
        client.post("/upload", base_url=BASE, data=data, content_type="multipart/form-data").status_code
        == 403
    )
    assert not any((app.config["WORKDIR"]).iterdir())  # nothing was saved


def test_security_headers(client):
    h = client.get("/", base_url=BASE).headers
    assert h["Cache-Control"] == "no-store" and h["X-Frame-Options"] == "DENY"
    csp = h["Content-Security-Policy"]
    assert "default-src 'none'" in csp and "http" not in csp and "unsafe" not in csp
    assert h["Referrer-Policy"] == "no-referrer"


def test_pages_load_no_external_resources(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    for url in ("/", f"/s/{sid}", "/restore"):
        html = client.get(url, base_url=BASE).get_data(as_text=True)
        assert (
            not re.search(r'(src|href)="https?://', html) and "<style" not in html and "onclick" not in html
        )


def test_unknown_session_and_bad_download_kind_are_404(client, app):
    assert client.get("/s/nonexistent", base_url=BASE).status_code == 404
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    for kind in ("..%2F..%2Fetc%2Fpasswd", "input", "mapping", "restored", "output"):
        assert client.get(f"/s/{sid}/download/{kind}", base_url=BASE).status_code == 404  # no result yet


def test_upload_size_limit(app, client):
    app.config["MAX_CONTENT_LENGTH"] = 200
    r = upload(client, app, b"x" * 5000, "big.csv")
    assert r.status_code == 413 and "limite" in r.get_data(as_text=True)


# -- upload and field list -----------------------------------------------------------------------


def test_fields_page_lists_fields_with_previews(client, app):
    r = upload(client, app, CSV, "clienti.csv")
    assert r.status_code == 302
    html = client.get(r.headers["Location"], base_url=BASE).get_data(as_text=True)
    for name in ("Nome", "Cognome", "Importo", "Citta", "Data"):
        assert f'<span class="name">{name}</span>' in html
    assert "1.234,56" in html and "Mario" in html
    assert 'value="" ' not in html  # no category is preselected: the app guesses nothing
    assert html.count("Scegli il tipo") == 5
    assert all(label in html for _, label in CATEGORIES)


def test_hostile_content_is_escaped(client, app):
    raw = b"<img src=x onerror=alert(1)>;b\n<script>alert(1)</script>;2\n"
    html = client.get(upload(client, app, raw, "x.csv").headers["Location"], base_url=BASE).get_data(
        as_text=True
    )
    assert "<script>alert" not in html and "<img src=x" not in html and "&lt;script&gt;" in html


def test_unsupported_and_invalid_uploads_show_a_clear_error_and_leave_nothing(client, app):
    for raw, name, expect in [
        (b"%PDF", "a.pdf", "non supportato"),
        (b"", "a.csv", "vuoto"),
        (b"{bad", "a.json", "JSON non valido"),
        (b"x", "a.xlsx", "non è un .xlsx"),
    ]:
        r = upload(client, app, raw, name)
        assert r.status_code == 400 and expect in r.get_data(as_text=True)
    r = client.post("/upload", base_url=BASE, data={"csrf": token(app)}, content_type="multipart/form-data")
    assert r.status_code == 400 and "Scegli un file" in r.get_data(as_text=True)
    assert not any(app.config["WORKDIR"].iterdir())


def test_filename_cannot_escape_the_session_folder(client, app):
    sid = sid_of(upload(client, app, CSV, "../../evil.csv"))
    sess = app.extensions["velo_sessions"][sid]
    assert sess.filename == "evil.csv" and (sess.dir / "in" / "evil.csv").exists()


# -- choosing and running ---------------------------------------------------------------------------


def test_validation_keeps_the_choices_and_creates_nothing(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    r = run(client, app, sid, {})
    assert r.status_code == 400 and "almeno un campo" in r.get_data(as_text=True)
    r = run(client, app, sid, {"sel_0": "on", "cat_0": "", "sel_1": "on", "cat_1": "last_name"})
    html = r.get_data(as_text=True)
    assert r.status_code == 400 and "«Nome»" in html
    assert re.search(r'id="sel1"[^>]*checked', html) and re.search(
        r'<option value="last_name"\s+selected', html
    )
    assert not list((app.extensions["velo_sessions"][sid].dir / "out").iterdir())


def test_full_csv_flow_pseudonymize_download_restore(client, app):
    sid = sid_of(upload(client, app, CSV, "clienti.csv"))
    r = run(client, app, sid, CSV_FORM)
    assert r.status_code == 302 and r.headers["Location"].endswith("/result")
    page = client.get(r.headers["Location"], base_url=BASE).get_data(as_text=True)
    assert "clienti.pseudo.csv" in page and "dati originali" in page and "Non condividerlo" in page

    out = client.get(f"/s/{sid}/download/output", base_url=BASE)
    assert "attachment" in out.headers["Content-Disposition"]
    body = out.data.decode("cp1252")
    assert "Mario" not in body and "Rossi" not in body and "Forlì" not in body
    assert "[CITY_" in body and "1.234,56" in body  # placeholders honoured, unchosen column intact

    mp = client.get(f"/s/{sid}/download/mapping", base_url=BASE)
    assert "attachment" in mp.headers["Content-Disposition"] and b"ATTENZIONE" in mp.data

    back = client.post(
        "/restore",
        base_url=BASE,
        content_type="multipart/form-data",
        data={
            "csrf": token(app),
            "pseudonymized": (io.BytesIO(out.data), "clienti.pseudo.csv"),
            "mapping": (io.BytesIO(mp.data), "clienti.pseudo.csv.velo-map.json"),
        },
    )
    assert back.status_code == 302
    rsid = sid_of(back)
    page = client.get(back.headers["Location"], base_url=BASE).get_data(as_text=True)
    assert "identico all'originale" in page and "clienti.restored.csv" in page
    assert client.get(f"/s/{rsid}/download/restored", base_url=BASE).data == CSV


def test_rerunning_with_other_choices_replaces_previous_output(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    run(client, app, sid, {"sel_0": "on", "cat_0": "first_name"})
    run(client, app, sid, {"sel_1": "on", "cat_1": "last_name"})
    body = client.get(f"/s/{sid}/download/output", base_url=BASE).data.decode("cp1252")
    assert "Mario" in body and "Rossi" not in body


def test_project_key_option_and_utf8_option(client, app, tmp_path, monkeypatch):
    monkeypatch.setenv("XDG_CONFIG_HOME", str(tmp_path / "cfg"))
    monkeypatch.setenv("APPDATA", str(tmp_path / "cfg"))
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    r = run(client, app, sid, {**CSV_FORM, "key": "project", "encoding": "utf-8"})
    page = client.get(r.headers["Location"], base_url=BASE).get_data(as_text=True)
    assert "salvata su questo computer" in page
    out = client.get(f"/s/{sid}/download/output", base_url=BASE).data
    out.decode("utf-8")  # valid UTF-8 now
    assert (tmp_path / "cfg" / "velo" / "project.key").exists()


def test_excel_flow(client, app, tmp_path):
    raw = make_workbook(tmp_path / "x.xlsx").read_bytes()
    r = upload(client, app, raw, "x.xlsx")
    sid = sid_of(r)
    html = client.get(r.headers["Location"], base_url=BASE).get_data(as_text=True)
    assert "foglio Clienti" in html and "foglio Archivio" in html and "nascosto" in html
    assert "3 celle con formule non vengono toccate" in html
    sess = app.extensions["velo_sessions"][sid]
    idx = [f.id for f in sess.fields].index("Clienti!Nome")
    run(client, app, sid, {f"sel_{idx}": "on", f"cat_{idx}": "first_name"})  # 'scrub' unchecked -> kept
    out = client.get(f"/s/{sid}/download/output", base_url=BASE).data
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb["Clienti"]["A2"].value != "Mario" and wb["Clienti"]["E2"].value == "=C2*0.22"
    assert wb.properties.creator == "Giuseppe Bianco"  # option honoured when unchecked


def test_json_flow(client, app):
    raw = b'{"clienti":[{"nome":"Mario Rossi","importo":1.10}]}'
    sid = sid_of(upload(client, app, raw, "d.json"))
    html = client.get(f"/s/{sid}", base_url=BASE).get_data(as_text=True)
    assert "clienti[].nome" in html and "clienti[].importo" in html
    run(client, app, sid, {"sel_0": "on", "cat_0": "fullname"})
    out = client.get(f"/s/{sid}/download/output", base_url=BASE).data
    assert b"Mario" not in out and b'"importo":1.10' in out


# -- restore errors and cleanup ----------------------------------------------------------------------


def test_restore_with_wrong_mapping_shows_an_error(client, app):
    r = client.post(
        "/restore",
        base_url=BASE,
        content_type="multipart/form-data",
        data={
            "csrf": token(app),
            "pseudonymized": (io.BytesIO(b"a;b\n1;2\n"), "x.csv"),
            "mapping": (io.BytesIO(b"{}"), "m.json"),
        },
    )
    assert r.status_code == 400 and "dizionario" in r.get_data(as_text=True)
    assert not any(app.config["WORKDIR"].iterdir())


def test_delete_removes_files_and_session(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    run(client, app, sid, CSV_FORM)
    folder = app.extensions["velo_sessions"][sid].dir
    assert folder.exists()
    r = client.post(f"/s/{sid}/delete", base_url=BASE, data={"csrf": token(app)})
    assert r.status_code == 200 and "eliminati" in r.get_data(as_text=True)
    assert not folder.exists()
    assert client.get(f"/s/{sid}", base_url=BASE).status_code == 404
    assert client.get(f"/s/{sid}/download/output", base_url=BASE).status_code == 404


def test_default_temp_folder_is_registered_for_cleanup(monkeypatch):
    import atexit

    calls = []
    monkeypatch.setattr(atexit, "register", lambda fn, *a, **k: calls.append((fn.__name__, a)))
    app = create_app(port=5000)
    assert calls and calls[0][0] == "rmtree" and calls[0][1][0] == app.config["WORKDIR"]
    import shutil

    shutil.rmtree(app.config["WORKDIR"])


def test_cli_args():
    assert parse_args([]).port == 8765 and parse_args(["--port", "9000", "--no-browser"]).no_browser


# -- cleanup of temporary files --------------------------------------------------------------------


def test_purge_stale_removes_only_old_velo_folders(tmp_path):
    import os
    import time

    from velo.web.app import purge_stale

    old, fresh, other = tmp_path / "velo-old", tmp_path / "velo-fresh", tmp_path / "altro-old"
    for d in (old, fresh, other):
        (d / "in").mkdir(parents=True)
        (d / "in" / "dati.csv").write_text("x")
    past = time.time() - 48 * 3600
    for d in (old, other):
        os.utime(d, (past, past))
    assert purge_stale(24, base=tmp_path) == 1
    assert not old.exists() and fresh.exists() and other.exists()


@pytest.mark.skipif(__import__("sys").platform == "win32", reason="POSIX signals")
def test_sigterm_removes_the_temporary_folder(tmp_path):
    import os
    import signal
    import socket
    import subprocess
    import sys
    import time

    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    env = {**os.environ, "TMPDIR": str(tmp_path), "TEMP": str(tmp_path), "TMP": str(tmp_path)}
    proc = subprocess.Popen(
        [sys.executable, "-m", "velo", "--no-browser", "--port", str(port)],
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        for _ in range(60):
            try:
                socket.create_connection(("127.0.0.1", port), timeout=0.2).close()
                break
            except OSError:
                time.sleep(0.1)
        assert list(tmp_path.glob("velo-*")), "the server should have created its folder"
        proc.send_signal(signal.SIGTERM)
        proc.wait(timeout=10)
    finally:
        if proc.poll() is None:
            proc.kill()
    assert not list(tmp_path.glob("velo-*"))


# -- consistency warning ------------------------------------------------------------------------------


def _totals_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Clienti"
    ws.append(["Nome", "IBAN", "Saldo"])
    ws.append(["Mario", "IT60X0542811101000000123456", 100])
    ws.append(["Lucia", "IT60X0542811101000000123456", 200])
    ws.append([None, "Totale", "=SUM(C2:C3)"])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    return p.read_bytes()


def _confirm_token(html):
    return re.search(r'name="confirmed" value="([0-9a-f]{16})"', html).group(1)


def test_mismatch_warning_then_explicit_confirmation(client, app, tmp_path):
    sid = sid_of(upload(client, app, _totals_workbook(tmp_path), "t.xlsx"))
    form = {"sel_1": "on", "cat_1": "iban", "mode_1": "placeholder", "sel_2": "on", "cat_2": "amount"}
    r = run(client, app, sid, form)
    html = r.get_data(as_text=True)
    assert r.status_code == 200 and "non corrispondono al tipo scelto" in html
    assert "1 valore su 3 non sembra un IBAN valido, per esempio «Totale»" in html
    assert "Sostituisci comunque" in html
    assert re.search(r'id="sel1"[^>]*checked', html)  # choices are kept
    assert not list((app.extensions["velo_sessions"][sid].dir / "out").iterdir())  # nothing written yet

    r = run(client, app, sid, {**form, "confirmed": _confirm_token(html)})
    assert r.status_code == 302
    page = client.get(r.headers["Location"], base_url=BASE).get_data(as_text=True)
    assert "Hai confermato valori fuori tipo" in page and "Totale" in page
    out = openpyxl.load_workbook(io.BytesIO(client.get(f"/s/{sid}/download/output", base_url=BASE).data))[
        "Clienti"
    ]
    assert out["B4"].value.startswith("[IBAN_") and out["C4"].value == "=SUM(C2:C3)"


def test_confirmation_does_not_survive_a_changed_choice(client, app, tmp_path):
    sid = sid_of(upload(client, app, _totals_workbook(tmp_path), "t.xlsx"))
    form = {"sel_1": "on", "cat_1": "iban"}
    token_ = _confirm_token(run(client, app, sid, form).get_data(as_text=True))
    other = {"sel_1": "on", "cat_1": "iban", "sel_2": "on", "cat_2": "date"}  # adds a new mismatch
    r = run(client, app, sid, {**other, "confirmed": token_})
    assert r.status_code == 200 and "non corrispondono" in r.get_data(as_text=True)
    assert run(client, app, sid, {"sel_1": "on", "cat_1": "iban", "confirmed": "0" * 16}).status_code == 200


def test_correcting_the_choice_removes_the_warning(client, app, tmp_path):
    sid = sid_of(upload(client, app, _totals_workbook(tmp_path), "t.xlsx"))
    r = run(client, app, sid, {"sel_0": "on", "cat_0": "first_name"})  # only the clean field
    assert r.status_code == 302


def test_no_confirmation_banner_when_data_is_clean(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    page = client.get(run(client, app, sid, CSV_FORM).headers["Location"], base_url=BASE).get_data(
        as_text=True
    )
    assert "confermato valori" not in page


# -- remaining edges ---------------------------------------------------------------------------------------


def test_file_without_fields_is_refused_and_cleaned_up(client, app):
    r = upload(client, app, b"{}", "empty.json")
    assert r.status_code == 400 and "non ci sono campi" in r.get_data(as_text=True)
    assert not any(app.config["WORKDIR"].iterdir())


def test_invalid_mode_value_falls_back_to_fake(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    run(client, app, sid, {"sel_0": "on", "cat_0": "first_name", "mode_0": "bogus"})
    body = client.get(f"/s/{sid}/download/output", base_url=BASE).data.decode("cp1252")
    assert "[FIRST_NAME_" not in body and "Mario" not in body


def test_service_errors_are_shown_on_the_form_with_choices_kept(client, app, monkeypatch):
    from velo.errors import EncodingError

    sid = sid_of(upload(client, app, CSV, "c.csv"))

    def boom(*a, **k):
        raise EncodingError("Il carattere 'Ł' non è rappresentabile")

    monkeypatch.setattr("velo.web.app.pseudonymize", boom)
    r = run(client, app, sid, CSV_FORM)
    html = r.get_data(as_text=True)
    assert (
        r.status_code == 400 and "non è rappresentabile" in html and re.search(r'id="sel0"[^>]*checked', html)
    )


def test_navigation_edges(client, app):
    sid = sid_of(upload(client, app, CSV, "c.csv"))
    r = client.get(f"/s/{sid}/result", base_url=BASE)  # nothing run yet
    assert r.status_code == 302 and r.headers["Location"].endswith(f"/s/{sid}")
    assert client.get(f"/s/{sid}/restored", base_url=BASE).status_code == 404
    assert client.get("/restore", base_url=BASE).status_code == 200


def test_error_pages_are_in_italian_and_leak_nothing(client):
    r = client.get("/nope", base_url=BASE)
    assert r.status_code == 404 and "non esiste" in r.get_data(as_text=True)
    assert "Traceback" not in r.get_data(as_text=True)


def test_main_reports_port_in_use_and_opens_browser_only_when_asked(monkeypatch):
    import velo.web.app as webapp

    class Stub:
        def __init__(self, error=None):
            self.error, self.ran = error, False

        def run(self, **kw):
            self.ran = True
            assert kw["host"] == "127.0.0.1" and kw["debug"] is False
            if self.error:
                raise self.error

    monkeypatch.setattr(webapp, "purge_stale", lambda: 0)
    timers = []
    monkeypatch.setattr(
        webapp.threading,
        "Timer",
        lambda delay, fn, args=(): timers.append((fn, args)) or type("T", (), {"start": lambda s: None})(),
    )

    stub = Stub()
    monkeypatch.setattr(webapp, "create_app", lambda port: stub)
    webapp.main(["--port", "9100"])
    assert stub.ran and timers and timers[0][1] == ("http://127.0.0.1:9100",)

    timers.clear()
    webapp.main(["--no-browser"])
    assert not timers

    monkeypatch.setattr(webapp, "create_app", lambda port: Stub(OSError("address in use")))
    with pytest.raises(SystemExit) as exc:
        webapp.main(["--no-browser", "--port", "9100"])
    assert "9100" in str(exc.value) and "--port" in str(exc.value)


def test_exit_handler_raises_system_exit_so_atexit_runs():
    from velo.web.app import _exit_cleanly

    with pytest.raises(SystemExit):
        _exit_cleanly(15, None)
